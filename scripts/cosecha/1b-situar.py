# -*- coding: utf-8 -*-
"""
Paso 1b: dar a cada municipio su ficha oficial y su punto en el mapa.

El mapa nuevo pinta municipios, no convocatorias, así que cada uno necesita una
coordenada. No se inventa ninguna: se toma de donde ya existe —nuestra propia
base y la API de porceuta.es— y se apunta de cuál de las dos, porque en este
proyecto ya ha habido coordenadas a 88 km de su municipio y sin saber de dónde
salieron no hay forma de rastrearlas.

Los que se queden sin punto no se esconden: salen en la lista para situarlos a
mano, y hasta entonces el mapa los enseña aparte.

    python3 scripts/cosecha/1b-situar.py [libro.xlsx]
"""
import json
import os
import re
import sys
import urllib.request

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "lib"))
import cruce  # noqa: E402
import rutas  # noqa: E402

import openpyxl  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
HOJA = os.path.join(rutas.datos(), "hoja.json")
PLACES = "https://ceutanosune.es/api/places"
PORCEUTA = "https://porceuta.es/api/public/convocatorias"
FECHA = "2026-09-02"


# Sin distinguir mayúsculas: el Registro escribe los artículos catalanes en
# minúscula («l'», «la», «el»), que es como van en catalán, y así se dejan.
ARTICULO_DETRAS = re.compile(r"^(.*?),\s*(El|La|Los|Las|L'|A|O|As|Os|Es|Sa|Els|Les)$", re.I)


def natural(nombre):
    """«Nava, La» → «La Nava»; «Hospitalet de Llobregat, L'» → «L'Hospitalet…».

    El Registro pone el artículo detrás para ordenar; en una ficha que lee una
    persona, va delante.
    """
    partes = []
    for trozo in str(nombre or "").split("/"):
        m = ARTICULO_DETRAS.match(trozo.strip())
        if m:
            art = m.group(2)
            trozo = f"{art}{'' if art.endswith(chr(39)) else ' '}{m.group(1)}"
        partes.append(trozo.strip())
    return "/".join(partes)


def bajar(url):
    pet = urllib.request.Request(url, headers={"user-agent": "ceuta-nos-une/cosecha"})
    with urllib.request.urlopen(pet, timeout=30) as r:
        return json.load(r)


def main():
    libro = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser(
        "~/Downloads/Municipios_España_Ayuntamientos_4.xlsx")
    hoja = json.load(open(HOJA, encoding="utf-8"))

    ws = openpyxl.load_workbook(libro, data_only=True)["Municipios España"]
    cab = [c.value for c in ws[1]]
    rel = [dict(zip(cab, fila)) for fila in ws.iter_rows(min_row=2, values_only=True)
           if fila[3]]
    ficha = cruce.hacer_indice(rel)
    print(f"Registro de Entidades Locales: {len(rel)} municipios")

    # Los puntos: primero los nuestros, que son los revisados a mano en /admin.
    puntos = {}
    try:
        for p in bajar(PLACES).get("places", []):
            if p.get("lat") is None:
                continue
            f, _ = ficha(p["city"], p["province"])
            puntos.setdefault(cruce.clave_de(p["city"], p["province"], f),
                              (p["lat"], p["lon"], "places"))
        print(f"ceutanosune.es: {len(puntos)} municipios con punto")
    except Exception as err:
        print("aviso: no se pudo leer ceutanosune.es —", err)

    try:
        antes = len(puntos)
        for c in bajar(PORCEUTA).get("convocatorias", []):
            if c.get("latitud") is None or (c.get("fecha") or "") != FECHA:
                continue
            f, _ = ficha(c["ciudad"], c["provincia"])
            puntos.setdefault(cruce.clave_de(c["ciudad"], c["provincia"], f),
                              (c["latitud"], c["longitud"], "porceuta"))
        print(f"porceuta.es: {len(puntos) - antes} municipios más")
    except Exception as err:
        print("aviso: no se pudo leer porceuta.es —", err)

    sin_punto, sin_ficha = [], []
    for m in hoja["municipios"]:
        f, via = ficha(m["municipio"], m["provincia"])
        m["clave"] = cruce.clave_de(m["municipio"], m["provincia"], f)
        m["ine"] = str(f["Código Municipio (REL)"]) if f and f.get("Código Municipio (REL)") else None
        m["poblacion"] = int(f["Población (habitantes)"]) if f and f.get("Población (habitantes)") else None
        m["via_rel"] = via
        # Para enseñar: el nombre en su forma natural y la provincia con el
        # nombre de siempre. Las claves de cruce no cambian.
        m["nombre"] = natural(m["municipio"])
        m["provincia_nombre"] = (cruce.provincia(f["Provincia"]) if f else None) \
            or cruce.provincia(m["provincia"]) or m["provincia"]
        if not f:
            sin_ficha.append(m["municipio"])
        punto = puntos.get(m["clave"])
        m["lat"], m["lon"], m["origen_pto"] = punto if punto else (None, None, None)
        if not punto:
            sin_punto.append(f'{m["municipio"]} ({m["provincia"]})')

    json.dump(hoja, open(HOJA, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    claves = {m["clave"] for m in hoja["municipios"]}
    print(f"\n{len(hoja['municipios'])} municipios · {len(claves)} claves distintas "
          f"(si no coinciden, dos filas son el mismo sitio y se fundirán al importar)")
    print(f"con punto en el mapa: {len(hoja['municipios']) - len(sin_punto)}")
    print(f"sin ficha en el REL: {len(sin_ficha)}")
    if sin_punto:
        print(f"\nsin punto ({len(sin_punto)}), para situarlos a mano:")
        for n in sin_punto[:40]:
            print("   ", n)
        if len(sin_punto) > 40:
            print(f"    … y {len(sin_punto) - 40} más")


if __name__ == "__main__":
    main()
